import io
from datetime import timedelta
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.orders.fulfilment import deliver, ship
from apps.orders.reports import RETURN_HEADERS, SALES_HEADERS, return_rows, sales_rows
from apps.orders.returns import inspect_return, open_return, refund_return
from apps.payments.models import Payment

from .conftest import NOW
from .test_fulfilment import imei, order, phone_line, units  # noqa: F401 — фикстуры
from .test_shipping import COURIER, picked

pytestmark = pytest.mark.django_db
TODAY = timezone.localdate(NOW)
REPORT = reverse("admin:orders_salesreport_changelist")


@pytest.fixture
def shipped(order):
    Payment.objects.create(order=order, provider="payme", external_id="t-1", amount_uzs=order.total_uzs,
                           state=Payment.State.PAID, created_at=NOW)
    return ship(picked(order), "own", COURIER, NOW, actor="manager:anna")


def test_sales_rows(shipped):
    phone, gift = sales_rows(TODAY, TODAY)
    row = dict(zip(SALES_HEADERS, phone))

    assert row["Заказ"] == shipped.number and row["Покупатель"] == "+998901234567"
    assert row["Товар"] == "iPhone 16 Pro · 256 · Чёрный" and row["Кол-во"] == 2
    assert row["Цена, сум"] == 14_080_000 and row["Сумма, сум"] == 28_160_000
    assert row["в т.ч. НДС, сум"] == 3_017_143  # 28 160 000 × 12 / 112
    assert row["Себестоимость, $"] == Decimal("2000.00") and row["Себестоимость, сум"] == 25_600_000
    assert row["IMEI"] == f"{imei(0)}, {imei(1)}" and row["Оплата"] == "Payme" and row["Подарок"] == ""
    gift_row = dict(zip(SALES_HEADERS, gift))
    assert (gift_row["Сумма, сум"], gift_row["в т.ч. НДС, сум"], gift_row["Подарок"]) == (0, 0, "да")


def test_period_is_inclusive_and_by_handover(shipped):
    assert sales_rows(TODAY - timedelta(days=5), TODAY - timedelta(days=1)) == []
    assert len(sales_rows(TODAY - timedelta(days=5), TODAY)) == 2


def test_return_rows(shipped, settings):
    settings.RETURN_DAYS = 10
    deliver(shipped, NOW, actor="manager:anna")
    open_return(shipped, "Не подошёл", NOW, actor="manager:anna")
    inspect_return(shipped, [imei(0), imei(1)], gift_returned=False, now=NOW, actor="manager:anna")
    refund_return(shipped, NOW, actor="manager:anna")

    [row] = return_rows(TODAY, TODAY)

    values = dict(zip(RETURN_HEADERS, row))
    assert values["Заказ"] == shipped.number and values["Сумма возврата, сум"] == 28_160_000 - 240_000
    assert values["Подарок вернули"] == "нет" and values["Причина"] == "Не подошёл"


def test_download(admin_client, shipped):
    response = admin_client.get(REPORT, {"date_from": TODAY.isoformat(), "date_to": TODAY.isoformat()})

    assert response["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert f"sales_{TODAY}_{TODAY}.xlsx" in response["Content-Disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Реализации", "Возвраты"]
    sales = book["Реализации"]
    assert [cell.value for cell in sales[1]] == SALES_HEADERS
    assert sales["B2"].value == shipped.number


@pytest.mark.parametrize("params", [{}, {"date_from": "вчера", "date_to": "2026-10-08"},
                                    {"date_from": "2026-10-09", "date_to": "2026-10-08"},
                                    {"date_from": "2024-01-01", "date_to": "2026-10-08"}])
def test_form_and_bad_periods(admin_client, params):
    response = admin_client.get(REPORT, params)

    assert response.status_code == 200 and "Выгрузка для бухгалтера" in response.content.decode()
    assert not response.has_header("Content-Disposition")


def test_report_needs_permission():
    staff = get_user_model().objects.create_user("guard", password="x", is_staff=True)
    client = Client()
    client.force_login(staff)

    assert client.get(REPORT, {"date_from": "2026-10-01", "date_to": "2026-10-08"}).status_code == 403
