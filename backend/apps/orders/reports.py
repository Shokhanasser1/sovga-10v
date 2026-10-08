"""Выгрузка для бухгалтера: реализации (по дате передачи в доставку) и возвраты (по дате возврата денег)."""

import io
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal

from django.conf import settings
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.catalog.labels import variant_title
from apps.delivery.models import Shipment
from apps.inventory.models import InventoryUnit
from apps.payments.models import Payment

from .models import Order, ReturnRequest

SALES_HEADERS = [
    "Дата реализации", "Заказ", "Покупатель", "Товар", "ИКПУ", "Код упаковки", "Кол-во", "Цена, сум", "Сумма, сум",
    "в т.ч. НДС, сум", "Подарок", "Себестоимость, $", "Курс", "Себестоимость, сум", "IMEI", "Оплата",
]
RETURN_HEADERS = ["Дата возврата денег", "Заказ", "Покупатель", "Сумма возврата, сум", "Подарок вернули", "IMEI", "Причина"]
MAX_DAYS = 366


def _bounds(date_from: date, date_to: date) -> tuple[datetime, datetime]:
    """Даты включительно, по ташкентскому времени."""
    zone = timezone.get_current_timezone()
    return (
        timezone.make_aware(datetime.combine(date_from, time.min), zone),
        timezone.make_aware(datetime.combine(date_to + timedelta(days=1), time.min), zone),
    )


def _whole(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _paid_with(order: Order) -> str:
    payment = order.payments.filter(state__in=[Payment.State.PAID, Payment.State.REFUNDED]).first()
    return payment.get_provider_display() if payment else order.get_payment_method_display()


def _imeis(units) -> str:
    return ", ".join(unit.imei1 or unit.serial for unit in units)


def sales_rows(date_from: date, date_to: date) -> list[list]:
    start, end = _bounds(date_from, date_to)
    vat = Decimal(settings.FISCAL_VAT_PERCENT)
    rows = []
    shipments = Shipment.objects.filter(handed_at__gte=start, handed_at__lt=end).select_related("order__customer")
    for shipment in shipments.order_by("handed_at", "pk"):
        order = shipment.order
        paid_with = _paid_with(order)
        for line in order.lines.select_related("variant__product_model__category").order_by("pk"):
            amount = line.price_uzs * line.qty
            cost_usd = line.cost_usd * line.qty
            rows.append([
                timezone.localdate(shipment.handed_at), order.number, order.customer.phone, variant_title(line.variant),
                line.variant.mxik_code, line.variant.package_code, line.qty, line.price_uzs, amount,
                _whole(Decimal(amount) * vat / (100 + vat)), "да" if line.is_gift else "", cost_usd, order.rate,
                _whole(cost_usd * order.rate), _imeis(line.units.order_by("pk")), paid_with,
            ])
    return rows


def return_rows(date_from: date, date_to: date) -> list[list]:
    start, end = _bounds(date_from, date_to)
    requests = ReturnRequest.objects.filter(
        status=ReturnRequest.Status.REFUNDED, refunded_at__gte=start, refunded_at__lt=end
    ).select_related("order__customer")
    return [
        [
            timezone.localdate(request.refunded_at), request.order.number, request.order.customer.phone,
            request.refund_amount_uzs, "да" if request.gift_returned else "нет",
            _imeis(InventoryUnit.objects.filter(order_line__order=request.order).order_by("pk")), request.reason,
        ]
        for request in requests.order_by("refunded_at", "pk")
    ]


def _sheet(book: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False) -> None:
    sheet = book.active if first else book.create_sheet()
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for index, header in enumerate(headers, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = max(12, len(header) + 2)
        for cell in sheet[letter][1:]:
            if isinstance(cell.value, date):
                cell.number_format = "DD.MM.YYYY"
            elif isinstance(cell.value, int) and "сум" in header:
                cell.number_format = "#,##0"
    sheet.freeze_panes = "A2"


def workbook(date_from: date, date_to: date) -> bytes:
    book = Workbook()
    _sheet(book, "Реализации", SALES_HEADERS, sales_rows(date_from, date_to), first=True)
    _sheet(book, "Возвраты", RETURN_HEADERS, return_rows(date_from, date_to))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
