"""Чтение прайсов поставщиков: CSV, Excel, JSON. Чистые функции, без БД."""

import csv
import io
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

REQUIRED_COLUMNS = ("code", "price", "qty")


@dataclass(frozen=True)
class FeedRow:
    code: str
    name: str
    price_usd: Decimal | None
    qty: int | None
    raw: dict = field(default_factory=dict)
    error: str = ""


def validate_columns(columns: Mapping) -> list[str]:
    return [f"не задана колонка: {key}" for key in REQUIRED_COLUMNS if not columns.get(key)]


def _plain(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_decimal(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        text = str(value)
    else:
        text = re.sub(r"[\s $]", "", str(value))
        text = text.replace(",", "") if "," in text and "." in text else text.replace(",", ".")
    try:
        number = Decimal(text)
        return number.quantize(Decimal("0.01")) if number.is_finite() else None
    except InvalidOperation:
        return None


def _to_int(value: object) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value >= 0 else None
    if isinstance(value, float):
        return int(value) if value.is_integer() and value >= 0 else None
    match = re.fullmatch(r">?\s*([0-9]+)\s*\+?", str(value).strip())
    return int(match.group(1)) if match else None


def parse_rows(records: Iterable[Mapping], columns: Mapping) -> list[FeedRow]:
    """Привести строки прайса к FeedRow; пустые строки пропускаются."""
    rows = []
    for record in records:
        raw = {str(key): _plain(value) for key, value in record.items()}
        if not any(raw.values()):
            continue
        code = raw.get(columns["code"], "")
        name = raw.get(columns.get("name") or "", "")
        price = _to_decimal(record.get(columns["price"]))
        qty = _to_int(record.get(columns["qty"]))
        if not code:
            error = "нет кода"
        elif price is None or price <= 0:
            error = "некорректная цена"
        elif qty is None:
            error = "некорректный остаток"
        else:
            error = ""
        rows.append(FeedRow(code, name, price, qty, raw, error))
    return rows


def read_csv(text: str) -> list[dict]:
    """CSV с разделителем «,», «;» или табуляцией — по заголовку."""
    text = text.lstrip("﻿")
    header = text.split("\n", 1)[0]
    delimiter = max([",", ";", "\t"], key=header.count)
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return [{key: value for key, value in row.items() if key is not None} for row in reader]


def read_xlsx(source) -> list[dict]:
    """Первый лист; заголовок — первая строка, где заполнено хотя бы две ячейки."""
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        header: list[str] | None = None
        records = []
        for values in workbook.active.iter_rows(values_only=True):
            filled = [value for value in values if value not in (None, "")]
            if header is None:
                if len(filled) >= 2:
                    header = [_plain(value) for value in values]
                continue
            if not filled:
                continue
            records.append({name: value for name, value in zip(header, values) if name})
        return records
    finally:
        workbook.close()


def read_json(payload: object) -> list[dict]:
    if isinstance(payload, dict):
        for key in ("items", "data", "products"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
    if isinstance(payload, list) and all(isinstance(item, dict) for item in payload):
        return payload
    raise ValueError('Ожидается список объектов или {"items": [...]}')
